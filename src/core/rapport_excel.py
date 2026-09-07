"""Génération du rapport Excel à partir des résultats de l'analyse."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def generer_rapport_excel(resultat: dict[str, Any], chemin_rapport: str) -> Path:
    """Crée un unique onglet avec la synthèse en haut et le détail par dossier dessous."""
    destination = Path(chemin_rapport)
    if destination.suffix.lower() != ".xlsx":
        destination = destination.with_suffix(".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)

    classeur = Workbook()
    rapport = classeur.active
    rapport.title = "Rapport"

    # La synthèse est placée avant le tableau détaillé, afin que le rapport soit
    # lisible d'un seul coup d'œil dans un unique onglet Excel.
    rapport.merge_cells("A1:E1")
    rapport["A1"] = "Rapport de vérification de présence"
    rapport["A1"].font = Font(bold=True, size=14)
    rapport["A3"] = "Dossier parent"
    rapport["B3"] = resultat["dossier_parent"]
    rapport["A4"] = "Sous-dossiers analysés"
    rapport["B4"] = resultat["nombre_dossiers"]
    rapport["A5"] = "Fichiers manquants"
    rapport["B5"] = resultat["nombre_fichiers_manquants"]
    for cellule in (rapport["A3"], rapport["A4"], rapport["A5"]):
        cellule.font = Font(bold=True)

    # Le tableau détaillé commence après une ligne vide qui sépare clairement les
    # informations générales des résultats à lire ou filtrer.
    ligne_entetes = 7
    entetes = ["Nom du dossier", "Chemin", "Nombre de fichiers manquants", "Fichiers manquants", "Erreur"]
    for colonne, entete in enumerate(entetes, start=1):
        rapport.cell(row=ligne_entetes, column=colonne, value=entete)
    couleur_entete = PatternFill("solid", fgColor="1F4E78")
    for cellule in rapport[ligne_entetes]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = couleur_entete

    # Une ligne par dossier permet de filtrer et trier facilement dans Excel.
    for item in resultat["resultats"]:
        rapport.append([
            item["nom_dossier"], item["chemin_dossier"], len(item["fichiers_manquants"]),
            ", ".join(item["fichiers_manquants"]), item["erreur"],
        ])
    rapport.freeze_panes = "A8"
    rapport.auto_filter.ref = f"A{ligne_entetes}:E{rapport.max_row}"
    for index, largeur in enumerate([28, 65, 28, 50, 50], start=1):
        rapport.column_dimensions[get_column_letter(index)].width = largeur

    classeur.save(destination)
    return destination
